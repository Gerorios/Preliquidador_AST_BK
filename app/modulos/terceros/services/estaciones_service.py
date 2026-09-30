"""Lo que facturó cada estación de servicio, y el cruce con lo que se cargó.

Etapa 10 del plan. Hasta acá el módulo sabía cuánto combustible se cargó según
el sistema de campo. Esto trae el otro lado: lo que la estación cobró.

**Una carga que la estación facturó y nadie registró es plata que la empresa
pagó y no le descontó a nadie.** Ése es el punto de toda la etapa.

El puente es el **número de vale**, y funciona: de los 14 vales de agosto de
YPF Oasis Alderete que cruzan, los 14 coinciden en litros al decimal.

─── Por qué el mapeo vive en la base y no en el código ──────────────────────

Cada estación manda un archivo distinto, y no un poco distinto:

    Calchaqui   .xlsx, 45 columnas, el vale en `NumVehiculo` mezclado con
                texto: '60277', '0332 ORDEN', 'ORDEN:60314', 'BIDON'
    Garsa       .xls de Excel 97, el vale en `ORDEN_CARGA`, la cantidad en
                NEGATIVO y la fecha como serial de Excel

Poner eso en el código obliga a un deploy cada vez que una estación cambia su
reporte. En una fila de `terceros_estacion` lo cambia quien lo ve romperse.

─── Lo que NO se filtra acá ─────────────────────────────────────────────────

Las líneas de nafta —la flota liviana, que queda afuera de la liquidación— se
guardan igual. Lo que no se guarda no se puede contar, y saber cuánto se fue en
flota liviana es una pregunta legítima aunque no se liquide a nadie.
"""
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from sqlalchemy.orm import Session

from app.modulos.terceros.models import (
    CargaCombustible, CargaFacturada, Estacion, Liquidacion,
)
from app.modulos.terceros.services import quincenas

# Un vale es un número de al menos cuatro dígitos. El campo viene con texto
# alrededor —'0332 ORDEN', 'ORDEN:60314'— y a veces sin número: 'BIDON',
# 'MONTANA'. Esos últimos son la flota liviana, que no lleva vale.
_NUMERO = re.compile(r"\d{4,}")

# Sólo letras y números: la patente llega 'FRN 606' de un lado y 'FRN606' del
# otro, y en el mismo archivo aparece de las dos formas.
_NO_ALFANUM = re.compile(r"[^A-Z0-9]")

# Qué productos son de la flota liviana. No se descartan: se marcan, para poder
# contarlos aparte sin que ensucien el cruce.
_NAFTA = ("NAFTA", "SUPER", "V-POWER", "PREMIUM")


class ArchivoInvalido(ValueError):
    """El archivo no se pudo leer con el mapeo de esa estación. El mensaje lo
    lee una persona, así que dice qué falta y no qué excepción saltó."""


def vale_de(valor) -> str | None:
    """El número de vale que hay adentro del campo, si lo hay."""
    m = _NUMERO.search(str(valor or ""))
    return m.group(0) if m else None


def es_flota_liviana(producto: str | None) -> bool:
    texto = str(producto or "").upper()
    return any(p in texto for p in _NAFTA)


def patente_de(valor) -> str | None:
    limpia = _NO_ALFANUM.sub("", str(valor or "").upper())
    return limpia or None


# Cuántos días de diferencia se toleran al cruzar sin vale. La estación fecha
# el remito y el sistema de campo fecha la carga, y casi nunca es el mismo día:
# en agosto, 21 de 23 cargas de Shell Famaillá que parecían no estar tenían su
# par un día antes, con la misma patente y los mismos litros.
DIAS_DE_GRACIA = 2


def _a_un_caracter(a: str, b: str) -> bool:
    """Si dos patentes difieren en una sola posición.

    La estación tipea la patente a mano y se equivoca: escribe IKI028 donde el
    colectivo es ILI028. Esa línea cae en «sin asignar» —esa patente no es de
    nadie— y su par real queda en «sin facturar»: **un error de tipeo infla las
    dos listas a la vez**.

    Un carácter y no una distancia de edición: agregar o sacar un carácter
    cambia el largo, y las patentes no varían de largo. Con el mismo largo, una
    posición distinta es un dedo errado; dos ya es otra patente.
    """
    if len(a) != len(b):
        return False
    return sum(1 for x, y in zip(a, b) if x != y) == 1


def _sin_fecha(patente, litros) -> tuple | None:
    """Qué identifica una carga cuando el vale no alcanza.

    Hace falta por dos estaciones distintas:

      - Una escribe **los últimos cuatro dígitos** del vale: donde el sistema
        de campo tiene 60457, el archivo dice '0457 ORDEN'. Truncar el número
        para que coincida sería adivinar; 60457 y 70457 darían igual.
      - Shell Famaillá, que es la más grande, **no manda el vale**: las 104
        líneas de agosto vienen con ese campo vacío.

    La patente y los litros sí son exactos. La fecha se compara aparte, con
    tolerancia, porque los dos sistemas fechan cosas distintas.
    """
    if litros is None:
        return None
    p = patente_de(patente)
    if not p:
        return None
    return (p, Decimal(str(litros)).quantize(Decimal("0.01")))


def _decimal(valor, absoluto: bool = False) -> Decimal | None:
    if valor is None or valor == "":
        return None
    try:
        d = Decimal(str(valor))
    except (InvalidOperation, TypeError, ValueError):
        return None
    return abs(d) if absoluto else d


def _fecha(valor, serial: bool, datemode: int = 0) -> date | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if serial:
        # Garsa manda la fecha como número de días desde 1900, que es como la
        # guarda Excel por dentro.
        try:
            import xlrd
            return xlrd.xldate_as_datetime(float(valor), datemode).date()
        except Exception:
            return None
    try:
        return date.fromisoformat(str(valor)[:10])
    except ValueError:
        return None


# ─── Leer el archivo ────────────────────────────────────────────────────────

def _filas_xlsx(datos: bytes, hoja: str, fila_encabezado: int):
    import io

    import openpyxl
    libro = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    if hoja not in libro.sheetnames:
        raise ArchivoInvalido(
            "El archivo no tiene una hoja «%s». Tiene: %s."
            % (hoja, ", ".join(libro.sheetnames)))
    ws = libro[hoja]
    it = ws.iter_rows(values_only=True)
    for _ in range(fila_encabezado - 1):
        next(it, None)
    encabezado = [str(c).strip() if c is not None else "" for c in next(it, [])]
    filas = [dict(zip(encabezado, f)) for f in it
             if f and any(v is not None for v in f)]
    libro.close()
    return filas, 0


def _filas_xls(datos: bytes, hoja: str, fila_encabezado: int):
    import xlrd
    libro = xlrd.open_workbook(file_contents=datos)
    if hoja not in libro.sheet_names():
        raise ArchivoInvalido(
            "El archivo no tiene una hoja «%s». Tiene: %s."
            % (hoja, ", ".join(libro.sheet_names())))
    ws = libro.sheet_by_name(hoja)
    n = fila_encabezado - 1
    encabezado = [str(ws.cell_value(n, c)).strip() for c in range(ws.ncols)]
    filas = [dict(zip(encabezado, [ws.cell_value(f, c) for c in range(ws.ncols)]))
             for f in range(n + 1, ws.nrows)]
    return filas, libro.datemode


def leer(datos: bytes, nombre: str, mapeo: dict) -> list[dict]:
    """Las líneas del archivo, ya traducidas a nuestros nombres.

    `mapeo` dice qué hoja, en qué fila está el encabezado, y qué columna del
    archivo es cada campo nuestro. Sale de `terceros_estacion.mapeo`.
    """
    if not mapeo:
        raise ArchivoInvalido(
            "Esta estación no tiene configurado cómo leer su archivo. "
            "Cargale el mapeo de columnas antes de subir uno.")

    hoja = mapeo.get("hoja", "Sheet")
    encabezado = int(mapeo.get("fila_encabezado", 1))
    lector = _filas_xls if nombre.lower().endswith(".xls") else _filas_xlsx
    filas, datemode = lector(datos, hoja, encabezado)

    columnas = mapeo.get("columnas", {})
    faltan = [c for c in ("vale", "litros") if c not in columnas]
    if faltan:
        raise ArchivoInvalido(
            "El mapeo de esta estación no dice qué columna es %s. Sin eso no "
            "se puede cruzar nada." % " ni ".join(faltan))

    serial = bool(mapeo.get("fecha_serial"))
    absolutos = set(mapeo.get("valor_absoluto", ()))

    def valor(fila, campo):
        col = columnas.get(campo)
        return fila.get(col) if col else None

    salida = []
    for fila in filas:
        linea = {
            "fecha": _fecha(valor(fila, "fecha"), serial, datemode),
            "vale": vale_de(valor(fila, "vale")),
            "litros": _decimal(valor(fila, "litros"), "litros" in absolutos),
        }
        # Sin fecha, sin vale y sin litros no es una carga: es el pie del
        # reporte o una fila en blanco que el Excel arrastra.
        if not any(linea.values()):
            continue
        producto = str(valor(fila, "producto") or "").strip() or None
        salida.append({
            **linea,
            "importe": _decimal(valor(fila, "importe"), "importe" in absolutos),
            "producto": producto,
            "patente": str(valor(fila, "patente") or "").strip() or None,
            "chofer": str(valor(fila, "chofer") or "").strip() or None,
        })
    return salida


# ─── Guardar y cruzar ───────────────────────────────────────────────────────

class EstacionesService:
    def __init__(self, db: Session):
        self.db = db

    def listar(self) -> list[Estacion]:
        return (self.db.query(Estacion)
                .order_by(Estacion.activa.desc(), Estacion.nombre).all())

    def lineas_por_estacion(self, quincena: date) -> dict[int, int]:
        """Cuántas líneas tiene cargadas cada estación en esa quincena.

        Es lo que dice si a una estación le falta subir el archivo, que es la
        pregunta con la que se entra a la pantalla.
        """
        from sqlalchemy import func
        filas = (self.db.query(CargaFacturada.estacion_id, func.count())
                 .filter(CargaFacturada.quincena == quincena)
                 .group_by(CargaFacturada.estacion_id).all())
        return {estacion_id: n for estacion_id, n in filas}

    def subir(self, estacion_id: int, quincena: date, datos: bytes,
              nombre: str, usuario_id: int | None = None) -> dict:
        """Guarda lo facturado por esa estación.

        **Cada línea va a la quincena de SU fecha**, no a la del archivo: el
        reporte de agosto de una de las estaciones arranca el 1 y termina el
        19, así que cae en las dos. `quincena` es sólo el respaldo para las
        líneas que vengan sin fecha.

        Reemplaza y no acumula: subir el mismo archivo dos veces es lo que pasa
        cuando alguien no está seguro de haberlo subido, y sumarlo dos veces
        duplicaría todo lo facturado sin que se note. Se reemplaza **cada
        quincena que el archivo toca**, no sólo la elegida.
        """
        estacion = self.db.query(Estacion).filter(Estacion.id == estacion_id).first()
        if estacion is None:
            raise ArchivoInvalido("No existe esa estación.")

        mapeo = estacion.mapeo
        if isinstance(mapeo, str):
            mapeo = json.loads(mapeo)
        lineas = leer(datos, nombre, mapeo)

        for l in lineas:
            l["quincena"] = (quincenas.inicio_de_quincena(l["fecha"])
                             if l["fecha"] else quincena)

        tocadas = sorted({l["quincena"] for l in lineas} | {quincena})
        borradas = (self.db.query(CargaFacturada)
                    .filter(CargaFacturada.estacion_id == estacion_id,
                            CargaFacturada.quincena.in_(tocadas))
                    .delete(synchronize_session=False))
        ahora = datetime.now()
        if lineas:
            self.db.bulk_insert_mappings(CargaFacturada, [
                {"estacion_id": estacion_id, "archivo": nombre,
                 "subido_en": ahora, "subido_por": usuario_id, **l} for l in lineas])
        self.db.commit()

        return {
            "estacion": estacion.nombre,
            "lineas": len(lineas),
            "reemplazadas": borradas,
            "quincenas": [q.isoformat() for q in tocadas],
            "sin_vale": sum(1 for l in lineas if not l["vale"]),
            "flota_liviana": sum(1 for l in lineas if es_flota_liviana(l["producto"])),
        }

    def agregar_a_mano(self, estacion_id: int, lineas: list[dict],
                       usuario_id: int | None = None) -> int:
        """Carga líneas tipeadas, sin archivo.

        Es para La Angostura, que manda los remitos por foto. A diferencia de
        subir un archivo, esto **suma**: se tipea de a poco, a medida que van
        llegando las fotos, y borrar lo anterior en cada carga sería perder lo
        que alguien acaba de escribir.

        La quincena sale de la fecha de cada línea, igual que en el archivo.
        """
        estacion = self.db.query(Estacion).filter(Estacion.id == estacion_id).first()
        if estacion is None:
            raise ArchivoInvalido("No existe esa estación.")

        ahora = datetime.now()
        filas = []
        for l in lineas:
            fecha = l.get("fecha")
            if fecha is None:
                raise ArchivoInvalido(
                    "Cada carga necesita su fecha: es lo que decide en qué "
                    "quincena entra.")
            filas.append({
                "estacion_id": estacion_id,
                "quincena": quincenas.inicio_de_quincena(fecha),
                "fecha": fecha,
                "vale": vale_de(l.get("vale")) or (str(l.get("vale") or "").strip() or None),
                "litros": _decimal(l.get("litros")),
                "importe": _decimal(l.get("importe")),
                "producto": (l.get("producto") or "").strip() or None,
                "patente": (l.get("patente") or "").strip() or None,
                "chofer": (l.get("chofer") or "").strip() or None,
                "archivo": None,      # nulo dice que se tipeó
                "subido_en": ahora,
                "subido_por": usuario_id,
            })
        if filas:
            self.db.bulk_insert_mappings(CargaFacturada, filas)
            self.db.commit()
        return len(filas)

    def borrar_linea(self, id_: int) -> None:
        """Sacar una línea tipeada mal. Las de archivo se corrigen volviendo a
        subirlo, pero una tipeada no tiene de dónde volver a salir."""
        fila = (self.db.query(CargaFacturada)
                .filter(CargaFacturada.id == id_).first())
        if fila is None:
            raise ArchivoInvalido("No existe esa línea.")
        self.db.delete(fila)
        self.db.commit()

    def lineas_de(self, estacion_id: int, quincena: date) -> list[CargaFacturada]:
        return (self.db.query(CargaFacturada)
                .filter(CargaFacturada.estacion_id == estacion_id,
                        CargaFacturada.quincena == quincena)
                .order_by(CargaFacturada.fecha, CargaFacturada.vale).all())

    def definir_origen(self, estacion_id: int, origen_campo: str | None):
        """Con qué nombre se registran sus cargas en el sistema de campo.

        Sin esto no hay nada que cruzar, y no se puede adivinar: el archivo se
        llama «Calchaqui» y en el sistema de campo esa estación es
        `YPF OASIS ALDERETE`.
        """
        estacion = self.db.query(Estacion).filter(Estacion.id == estacion_id).first()
        if estacion is None:
            raise ArchivoInvalido("No existe esa estación.")
        estacion.origen_campo = (origen_campo or "").strip() or None
        self.db.commit()
        return estacion

    # ─── El cruce ───────────────────────────────────────────────────────────

    def vales_facturados(self, quincena: date) -> dict[str, CargaFacturada]:
        """Por vale, lo que la estación facturó en esa quincena.

        Es lo que la grilla usa para marcar cada carga. Un vale repetido en el
        archivo se queda con el primero: que la estación facture dos veces el
        mismo vale es otro problema, y lo cuenta `cruce`.
        """
        salida: dict[str, CargaFacturada] = {}
        for f in (self.db.query(CargaFacturada)
                  .filter(CargaFacturada.quincena == quincena,
                          CargaFacturada.vale.isnot(None)).all()):
            salida.setdefault(f.vale, f)
        return salida

    def cruce(self, quincena: date, patentes_colectivos: set[str] | None = None) -> dict:
        """Qué encontró y qué no encontró cada lado.

        Las dos preguntas que importan: qué facturó la estación que nadie
        cargó —plata pagada y no descontada— y dónde los litros no coinciden.

        `patentes_colectivos` son las patentes del maestro del sistema de
        campo. Sin eso, la lista de "no cargado" se llena de cosas que nunca
        van a estar: la cisterna, los bidones, las camionetas y la maquinaria
        no se cargan como flete y no le corresponden a ningún tercero. Una
        alerta que avisa veinte cosas de las que dieciocho son normales se
        deja de leer a la semana.
        """
        facturadas = [f for f in self.db.query(CargaFacturada)
                      .filter(CargaFacturada.quincena == quincena).all()]
        nuestras = self._cargas_de(quincena)
        por_vale: dict[str, list[CargaCombustible]] = {}
        for c in nuestras:
            v = (c.vale or "").strip()
            if v:
                por_vale.setdefault(v, []).append(c)

        # El segundo camino: misma patente y mismos litros, con la fecha cerca.
        por_patente_litros: dict[tuple, list[CargaCombustible]] = {}
        for c in nuestras:
            h = _sin_fecha(c.colectivo_patente, c.litros)
            if h:
                por_patente_litros.setdefault(h, []).append(c)

        # Una carga nuestra se usa una sola vez: si no, dos líneas facturadas
        # iguales cruzarían las dos contra la misma y una de las dos estaría
        # facturada de más sin que se vea.
        usadas: set[int] = set()

        filas = []
        por_vale_n, por_huella_n, por_parecida_n = 0, 0, 0
        for f in facturadas:
            nuestra, como = None, None
            for c in (por_vale.get(f.vale) or []):
                if c.id not in usadas:
                    nuestra = c
                    break
            if nuestra is not None:
                por_vale_n += 1
                como = "vale"
            else:
                nuestra = self._mas_cercana(
                    por_patente_litros.get(_sin_fecha(f.patente, f.litros)),
                    f.fecha, usadas)
                if nuestra is not None:
                    por_huella_n += 1
                    como = "patente_litros"
                elif not self._hay_que_avisar(f, patentes_colectivos):
                    # Sólo acá: si la patente SÍ es de un colectivo, no está
                    # tipeada mal y buscarle una parecida sería inventarle un
                    # dueño. LFL019 y LFL029 son dos colectivos distintos del
                    # mismo tercero, y están a un carácter uno del otro.
                    nuestra = self._por_patente_parecida(
                        f, por_patente_litros, usadas)
                    if nuestra is not None:
                        por_parecida_n += 1
                        como = "patente_parecida"
            if nuestra is not None:
                usadas.add(nuestra.id)

            if nuestra is None:
                como = ("sin_cargar" if self._hay_que_avisar(f, patentes_colectivos)
                        else "sin_asignar")
            elif (f.litros is not None and nuestra.litros is not None
                    and abs(Decimal(str(nuestra.litros)) - f.litros) > Decimal("0.5")):
                como = "litros_distintos"
            filas.append({"estado": como, "facturada": f, "carga": nuestra,
                          "estacion_id": f.estacion_id})

        # El otro lado: lo que el sistema de campo tiene y la estación no
        # facturó. Sólo se mira de las estaciones que subieron archivo — de las
        # que no, TODAS sus cargas aparecerían como faltantes y no faltan: lo
        # que falta es el archivo.
        con_archivo = self._origenes_con_archivo(quincena)
        cubre = self._dias_que_cubre(quincena)
        for c in nuestras:
            if c.id in usadas:
                continue
            estacion_id = con_archivo.get((c.origen or "").strip().upper())
            if estacion_id is None:
                continue
            filas.append({
                "estado": ("fuera_del_archivo"
                           if self._fuera_del_archivo(c, cubre.get(estacion_id))
                           else "sin_estacion"),
                "facturada": None, "carga": c, "estacion_id": estacion_id})

        return {
            "facturadas": len(facturadas),
            "cargadas": len(nuestras),
            "cruzan": por_vale_n + por_huella_n + por_parecida_n,
            "por_vale": por_vale_n,
            "por_huella": por_huella_n,
            "por_parecida": por_parecida_n,
            "filas": filas,
        }

    @staticmethod
    def _fuera_del_archivo(carga, rango) -> bool:
        """Si esta carga cae en días que el archivo de esa estación no cubre.

        Un reporte puede traer sólo dos días de la quincena. Todo lo que se
        cargó los otros trece aparecería como «la estación no lo facturó», y no
        es eso: es que no lo mandó. En agosto eran 6 de 19 avisos, todos de una
        estación cuyo archivo traía el 18 y el 19 y nada más.

        Sin fecha no se puede decidir, y la duda no se cuenta como error.
        """
        if rango is None or carga.fecha_uso is None:
            return False
        return not (rango[0] <= carga.fecha_uso <= rango[1])

    def _dias_que_cubre(self, quincena: date) -> dict[int, tuple]:
        """El primer y el último día que trae el archivo de cada estación."""
        from sqlalchemy import func
        filas = (self.db.query(CargaFacturada.estacion_id,
                               func.min(CargaFacturada.fecha),
                               func.max(CargaFacturada.fecha))
                 .filter(CargaFacturada.quincena == quincena,
                         CargaFacturada.fecha.isnot(None))
                 .group_by(CargaFacturada.estacion_id).all())
        return {e: (desde, hasta) for e, desde, hasta in filas}

    def _por_patente_parecida(self, facturada, indice, usadas):
        """La carga de una patente a un carácter de la que trae el archivo.

        Pide los mismos litros y la fecha cerca, igual que el cruce por
        patente: una patente parecida sola no alcanza para afirmar nada.

        Si hay más de una candidata no elige ninguna. Adivinar cuál de dos es
        peor que dejar el aviso: el aviso se mira, la adivinada no.
        """
        p = patente_de(facturada.patente)
        if not p or facturada.litros is None:
            return None
        huella = _sin_fecha(facturada.patente, facturada.litros)
        if huella is None:
            return None
        candidatas = [c for (otra, litros), cargas in indice.items()
                      if litros == huella[1] and _a_un_caracter(p, otra)
                      for c in cargas]
        elegidas = [c for c in candidatas
                    if self._mas_cercana([c], facturada.fecha, usadas) is not None]
        return elegidas[0] if len(elegidas) == 1 else None

    def _origenes_con_archivo(self, quincena: date) -> dict[str, int]:
        """Con qué nombre se registra en el sistema de campo cada estación que
        ya subió su archivo, y de qué estación es ese nombre.

        Es lo que permite distinguir «esta carga la estación no la facturó» de
        «todavía no subimos el archivo de esa estación». Sin esto, la segunda
        se leería como la primera y se reclamarían cargas que están bien.
        """
        con_lineas = {e for (e,) in self.db.query(CargaFacturada.estacion_id)
                      .filter(CargaFacturada.quincena == quincena).distinct()}
        if not con_lineas:
            return {}
        return {(e.origen_campo or "").strip().upper(): e.id
                for e in self.db.query(Estacion)
                .filter(Estacion.id.in_(con_lineas)).all()
                if e.origen_campo}

    def _mas_cercana(self, candidatas, fecha, usadas):
        """De las cargas con la misma patente y los mismos litros, la del día
        más cercano que todavía no se haya usado.

        La más cercana y no cualquiera: si un colectivo cargó los mismos litros
        dos veces en la semana, emparejarlas al revés daría el mismo total pero
        dejaría dos fechas cruzadas, y la próxima vez que alguien mire una
        línea no va a entender por qué.
        """
        if not candidatas or fecha is None:
            return None
        libres = [(abs((c.fecha_uso - fecha).days), c) for c in candidatas
                  if c.id not in usadas and c.fecha_uso is not None]
        libres = [(d, c) for d, c in libres if d <= DIAS_DE_GRACIA]
        if not libres:
            return None
        libres.sort(key=lambda x: x[0])
        return libres[0][1]

    def _hay_que_avisar(self, facturada, patentes_colectivos) -> bool:
        """Si esta línea facturada que no cruzó es algo que alguien tiene que
        ir a cargar, o algo que nunca iba a estar.

        No van a estar nunca, y no son un error:
          - la nafta, que es la flota liviana
          - lo que no salió a un colectivo: la cisterna, los bidones, la
            maquinaria, las camionetas
        """
        if es_flota_liviana(facturada.producto):
            return False
        if patentes_colectivos is None:
            return True
        p = patente_de(facturada.patente)
        return bool(p) and p in patentes_colectivos

    def _cargas_de(self, quincena: date) -> list[CargaCombustible]:
        from sqlalchemy import func
        quincena_liq = func.coalesce(CargaCombustible.quincena_efectiva,
                                     Liquidacion.quincena)
        return (self.db.query(CargaCombustible)
                .join(Liquidacion, Liquidacion.id == CargaCombustible.liquidacion_id)
                .filter(quincena_liq == quincena).all())
