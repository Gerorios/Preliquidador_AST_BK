"""Horas de taller sobre maquinaria de terceros, leídas de la app del taller.

Cuarta de las consultas de la etapa 1 (docs/modulos/terceros/plan-terceros.md).
A diferencia de las otras tres no sale de una base: la app del taller publica un
Google Sheet y el módulo lo lee solo, sin que nadie suba nada. La URL viene por
`TALLER_SHEET_URL` y no está en el repo (los repos son públicos y el Sheet expone
datos de los mecánicos).

Reproduce la consulta Power Query que hoy alimenta la hoja `App-Horas Taller` del
Excel (ver docs/modulos/terceros/fuentes/), acotada a una Quincena:

  - Del Sheet salen dos hojas: `Maestro_Maquinas` (qué máquina es de quién) y
    `BD_Horas` (las horas cargadas).
  - Se queda solo con las máquinas cuya propiedad es TERCEROS.
  - Deja afuera las horas RECHAZADAS. Las aprobadas y las pendientes vienen las
    dos, con su `estado`: cobrar solo las aprobadas es decisión del cálculo, no
    de la lectura (ver CONTEXT-terceros.md, "Hora de taller").
  - `horas_total` = horas + preparación + traslado.
  - El Tercero se deduce del nombre de la máquina, porque el Sheet no lo trae
    como campo propio (ver `_detectar_tercero`).

La descarga y el parseo están separados a propósito: así la lógica se testea con
un libro armado a mano, sin salir a la red.
"""
import io
from datetime import date, datetime

import httpx
import openpyxl

from app.core.config import settings
from app.core.quincena import calcular_rango_quincena
from app.modulos.terceros.services import quincenas

HOJA_MAESTRO = "Maestro_Maquinas"
HOJA_HORAS = "BD_Horas"

PROPIEDAD_TERCEROS = "TERCEROS"
ESTADO_APROBADO = "APROBADO"
ESTADO_PENDIENTE = "PENDIENTE"
ESTADO_RECHAZADO = "RECHAZADO"

# Columnas que se leen de cada hoja, con el nombre tal como lo escribe el Sheet.
COLUMNAS_MAESTRO = ("id_maquina", "nombre_maquinaria", "Tipo", "propiedad")
COLUMNAS_HORAS = (
    "Fecha", "Rubro", "Sub Rubro", "Horas", "Lugar", "Estado",
    "Horas_Preparacion", "Horas_Traslado", "Finca", "id_maquina",
)


class TallerNoConfigurado(RuntimeError):
    """Falta TALLER_SHEET_URL. Se avisa; no se inventa un origen."""


def descargar_libro(url: str | None = None, timeout: float = 60.0) -> bytes:
    """Baja el Sheet publicado como xlsx. Es la única parte que toca la red."""
    url = url or settings.taller_sheet_url
    if not url:
        raise TallerNoConfigurado(
            "Falta TALLER_SHEET_URL: sin ella no se pueden leer las horas de taller."
        )
    respuesta = httpx.get(url, timeout=timeout, follow_redirects=True)
    respuesta.raise_for_status()
    return respuesta.content


def _numero(valor) -> float:
    """Una celda vacía es 0 hora, no un error: así lo hace la consulta de hoy."""
    if valor is None or valor == "":
        return 0.0
    return float(valor)


def _id_maquina(valor) -> int | None:
    """El id viene como texto en el Sheet publicado y como número en una copia
    bajada a mano. Se lleva a entero de los dos lados para que el cruce no falle."""
    if valor is None or valor == "":
        return None
    return int(round(float(valor)))


def _fecha(valor) -> date | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return datetime.fromisoformat(str(valor)).date()


def _detectar_tercero(nombre: str | None) -> str:
    """El nombre del Tercero sale del nombre de la máquina: lo que viene después
    del número de interno, o después de la palabra MAQUINARIA. Se calcula acá y
    no se carga a mano, así una máquina nueva ya viene con su tercero.

    Es una heurística sobre texto, heredada de la consulta de hoy, y por eso el
    plan pide que el sistema de campo traiga el dueño como campo propio."""
    s = (nombre or "").strip()
    if not s:
        return s
    marcas = [
        t for t in s.split(" ")
        if t.upper().startswith("N") and any(c.isdigit() for c in t)
    ]
    if marcas:
        ultima = marcas[-1]
        corte = s[s.rindex(ultima) + len(ultima):].strip()
    elif s.upper().startswith("MAQUINARIA "):
        corte = s[len("MAQUINARIA "):].strip()
    else:
        corte = ""
    return corte or s




def _filas(hoja, columnas: tuple[str, ...]) -> list[dict]:
    """Lee una hoja por nombre de columna, no por posición: la app del taller
    puede agregar columnas y no queremos que eso corra los datos de lugar."""
    it = hoja.iter_rows(values_only=True)
    try:
        cabecera = next(it)
    except StopIteration:
        return []
    indice = {str(c).strip(): n for n, c in enumerate(cabecera) if c is not None}
    faltan = [c for c in columnas if c not in indice]
    if faltan:
        raise ValueError(
            f"La hoja '{hoja.title}' no tiene las columnas {faltan}. "
            "Si la app del taller cambió el formato, hay que ajustar la consulta."
        )
    return [
        {c: fila[indice[c]] for c in columnas}
        for fila in it
        if any(v is not None for v in fila)
    ]


def _leer_crudo(libro: bytes, quincena: date | None = None) -> list[dict]:
    """Las horas de la quincena sobre máquinas de terceros, **con las
    rechazadas incluidas**.

    Es la base de las dos vistas que el módulo necesita y que difieren sólo en
    eso: el listado de lo que se cobra (que las deja afuera, porque una hora
    rechazada no se cobra nunca) y el tablero de estados (que las cuenta, para
    poder reclamarle al taller antes de liquidar). Una sola pasada: si el
    cruce con el maestro o el recorte de la quincena se hicieran dos veces,
    las dos vistas podrían dejar de hablar del mismo conjunto.
    """
    wb = openpyxl.load_workbook(io.BytesIO(libro), read_only=True, data_only=True)
    try:
        maestro = _filas(wb[HOJA_MAESTRO], COLUMNAS_MAESTRO)
        horas = _filas(wb[HOJA_HORAS], COLUMNAS_HORAS)
    finally:
        wb.close()

    terceros = {
        _id_maquina(m["id_maquina"]): m
        for m in maestro
        if str(m["propiedad"] or "").strip().upper() == PROPIEDAD_TERCEROS
        and _id_maquina(m["id_maquina"]) is not None
    }

    desde, hasta = calcular_rango_quincena(quincena) if quincena else (None, None)

    resultado = []
    for h in horas:
        maquina = terceros.get(_id_maquina(h["id_maquina"]))
        if maquina is None:          # no es de un tercero: no se cobra
            continue
        fecha = _fecha(h["Fecha"])
        if fecha is None:            # sin fecha no hay quincena a la que imputar
            continue
        if desde and not (desde <= fecha <= hasta):
            continue
        nombre_maquina = maquina["nombre_maquinaria"]
        resultado.append({
            "fecha": fecha,
            "quincena_mes": quincenas.etiqueta(fecha),
            "anio": fecha.year,
            "tercero": _detectar_tercero(nombre_maquina),
            "maquina": nombre_maquina,
            "tipo_maquina": maquina["Tipo"],
            "rubro": h["Rubro"],
            "sub_rubro": h["Sub Rubro"],
            "finca": h["Finca"],
            "lugar": h["Lugar"],
            "estado": h["Estado"],
            "horas": _numero(h["Horas"]),
            "horas_preparacion": _numero(h["Horas_Preparacion"]),
            "horas_traslado": _numero(h["Horas_Traslado"]),
            "horas_total": (
                _numero(h["Horas"])
                + _numero(h["Horas_Preparacion"])
                + _numero(h["Horas_Traslado"])
            ),
            "id_maquina": _id_maquina(h["id_maquina"]),
        })

    resultado.sort(key=lambda r: (r["fecha"], r["tercero"]))
    return resultado


def _estado(fila: dict) -> str:
    return str(fila["estado"] or "").strip().upper()


def leer_horas(libro: bytes, quincena: date | None = None) -> list[dict]:
    """Las horas que se pueden cobrar. Sin red. `quincena` None devuelve todo."""
    return [h for h in _leer_crudo(libro, quincena) if _estado(h) != ESTADO_RECHAZADO]


def contar_estados(libro: bytes, quincena: date) -> dict:
    """Cuántas horas hay en cada estado, para el tablero de la quincena.

    Cuenta las rechazadas, que el listado no muestra: la pregunta que responde
    no es "qué cobro" sino "qué falta que el taller resuelva antes de que yo
    liquide" (plan-terceros.md, sección 2.4).
    """
    filas = _leer_crudo(libro, quincena)
    def horas_de(estado):
        return sum(f["horas_total"] for f in filas if _estado(f) == estado)
    return {
        "aprobadas": sum(1 for f in filas if _estado(f) == ESTADO_APROBADO),
        "pendientes": sum(1 for f in filas if _estado(f) == ESTADO_PENDIENTE),
        "rechazadas": sum(1 for f in filas if _estado(f) == ESTADO_RECHAZADO),
        "horas_aprobadas": horas_de(ESTADO_APROBADO),
        "horas_pendientes": horas_de(ESTADO_PENDIENTE),
    }


def leer_maestro(libro: bytes) -> list[dict]:
    """El maestro de máquinas del Sheet, para las Alertas de cruce.

    Trae las de terceros con su id, que es el que debería coincidir con el del
    sistema de compras. No filtra por quincena: un problema de cruce no es de
    una quincena, es del maestro.
    """
    wb = openpyxl.load_workbook(io.BytesIO(libro), read_only=True, data_only=True)
    try:
        filas = _filas(wb[HOJA_MAESTRO], COLUMNAS_MAESTRO)
    finally:
        wb.close()
    return [
        {
            "id_maquina": _id_maquina(f["id_maquina"]),
            "nombre": f["nombre_maquinaria"],
            "tipo": f["Tipo"],
            "propiedad": str(f["propiedad"] or "").strip().upper(),
        }
        for f in filas
        if str(f["propiedad"] or "").strip().upper() == PROPIEDAD_TERCEROS
        and _id_maquina(f["id_maquina"]) is not None
    ]


class ConsultaTallerService:
    """Las horas de taller de una quincena. `libro` permite pasar un xlsx ya
    bajado —un test, o una lectura que se reusa— en vez de salir a la red."""

    def __init__(self, libro: bytes | None = None):
        self._libro = libro

    def _bajar(self) -> bytes:
        """Se baja una sola vez por instancia: las dos vistas de una misma
        pantalla no tienen por qué pedir el Sheet dos veces."""
        if self._libro is None:
            self._libro = descargar_libro()
        return self._libro

    def horas_quincena(self, quincena: date) -> list[dict]:
        return leer_horas(self._bajar(), quincena)

    def estados_quincena(self, quincena: date) -> dict:
        return contar_estados(self._bajar(), quincena)

    def maestro(self) -> list[dict]:
        return leer_maestro(self._bajar())
