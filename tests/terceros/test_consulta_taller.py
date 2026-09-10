"""Las horas de taller leídas del Sheet de la app del taller (etapa 1).

Acá sí se prueba la lógica completa, porque es Python y no SQL: el libro se
arma en memoria y no se sale a la red.
"""
import io
from datetime import date

import openpyxl
import pytest

from app.modulos.terceros.services import consulta_taller
from app.modulos.terceros.services.consulta_taller import (
    ConsultaTallerService,
    TallerNoConfigurado,
    _detectar_tercero,
    descargar_libro,
    leer_horas,
)

MAESTRO = [
    # id_maquina, nombre_maquinaria, Tipo, propiedad
    (52, "TRACTOR A55 DEUTZ N°0001 TATO", "TRACTORES", "TERCEROS"),
    (735, "CARGADORA MANITOU N°9872 PABLO ROJAS", "CARGADORAS", "TERCEROS"),
    (982, "CINCEL SAN MIGUEL", "MAQUINAS GENERALES", "TERCEROS"),
    (100, "TRACTOR JOHN DEERE N°0500 LA ASTURIANA", "TRACTORES", "PROPIA"),
]


def _libro(filas_horas, maestro=None):
    """Un Sheet como el que publica la app del taller."""
    wb = openpyxl.Workbook()
    hm = wb.active
    hm.title = consulta_taller.HOJA_MAESTRO
    hm.append(consulta_taller.COLUMNAS_MAESTRO)
    for fila in (maestro if maestro is not None else MAESTRO):
        hm.append(fila)

    hh = wb.create_sheet(consulta_taller.HOJA_HORAS)
    hh.append(consulta_taller.COLUMNAS_HORAS)
    for fila in filas_horas:
        hh.append(fila)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _hora(fecha, id_maquina, estado="Aprobado", horas=1, prep=0, traslado=0):
    # Fecha, Rubro, Sub Rubro, Horas, Lugar, Estado, Horas_Preparacion,
    # Horas_Traslado, Finca, id_maquina
    return (fecha, "HERRERIA", "385. PERNO", horas, "TALLER", estado,
            prep, traslado, "EL CORTE", id_maquina)


# ─── Qué horas entran ───────────────────────────────────────────────────────

def test_solo_las_maquinas_de_terceros():
    """La máquina propia no se le cobra a nadie."""
    libro = _libro([_hora(date(2026, 8, 3), 52), _hora(date(2026, 8, 3), 100)])
    filas = leer_horas(libro, date(2026, 8, 1))
    assert [f["id_maquina"] for f in filas] == [52]


def test_las_rechazadas_no_se_cobran_nunca():
    libro = _libro([
        _hora(date(2026, 8, 3), 52, estado="Aprobado"),
        _hora(date(2026, 8, 4), 735, estado="Rechazado"),
        _hora(date(2026, 8, 5), 982, estado="RECHAZADO"),
    ])
    filas = leer_horas(libro, date(2026, 8, 1))
    assert [f["id_maquina"] for f in filas] == [52]


def test_las_pendientes_vienen_con_su_estado():
    """La lectura las trae; cobrar solo las aprobadas lo decide el cálculo, no
    esta consulta (ver CONTEXT-terceros.md, "Hora de taller")."""
    libro = _libro([
        _hora(date(2026, 8, 3), 52, estado="Aprobado"),
        _hora(date(2026, 8, 4), 735, estado="Pendiente"),
    ])
    filas = leer_horas(libro, date(2026, 8, 1))
    assert [f["estado"] for f in filas] == ["Aprobado", "Pendiente"]


def test_una_hora_sin_fecha_no_tiene_quincena_a_la_que_imputarse():
    libro = _libro([_hora(None, 52), _hora(date(2026, 8, 3), 735)])
    filas = leer_horas(libro, date(2026, 8, 1))
    assert [f["id_maquina"] for f in filas] == [735]


def test_una_maquina_que_el_maestro_no_conoce_queda_afuera():
    libro = _libro([_hora(date(2026, 8, 3), 999)])
    assert leer_horas(libro, date(2026, 8, 1)) == []


# ─── La quincena ────────────────────────────────────────────────────────────

def test_la_primera_quincena_toma_del_1_al_15():
    libro = _libro([
        _hora(date(2026, 7, 31), 52),
        _hora(date(2026, 8, 1), 52),
        _hora(date(2026, 8, 15), 735),
        _hora(date(2026, 8, 16), 982),
    ])
    filas = leer_horas(libro, date(2026, 8, 1))
    assert [f["fecha"] for f in filas] == [date(2026, 8, 1), date(2026, 8, 15)]


def test_la_segunda_quincena_toma_del_16_a_fin_de_mes():
    libro = _libro([
        _hora(date(2026, 8, 15), 52),
        _hora(date(2026, 8, 16), 52),
        _hora(date(2026, 8, 31), 735),
        _hora(date(2026, 9, 1), 982),
    ])
    filas = leer_horas(libro, date(2026, 8, 16))
    assert [f["fecha"] for f in filas] == [date(2026, 8, 16), date(2026, 8, 31)]


def test_sin_quincena_devuelve_todo_el_sheet():
    libro = _libro([_hora(date(2026, 3, 3), 52), _hora(date(2026, 8, 3), 735)])
    assert len(leer_horas(libro)) == 2


def test_la_quincena_del_mes_se_escribe_como_en_el_excel():
    libro = _libro([_hora(date(2026, 8, 3), 52), _hora(date(2026, 8, 20), 735)])
    filas = leer_horas(libro)
    assert [f["quincena_mes"] for f in filas] == ["08-1Q", "08-2Q"]
    assert [f["anio"] for f in filas] == [2026, 2026]


# ─── Las horas ──────────────────────────────────────────────────────────────

def test_la_hora_total_incluye_preparacion_y_traslado():
    libro = _libro([_hora(date(2026, 8, 3), 735, horas=2, prep=1, traslado=2)])
    fila = leer_horas(libro, date(2026, 8, 1))[0]
    assert (fila["horas"], fila["horas_preparacion"], fila["horas_traslado"]) == (2.0, 1.0, 2.0)
    assert fila["horas_total"] == 5.0


def test_las_celdas_vacias_de_horas_cuentan_como_cero():
    libro = _libro([_hora(date(2026, 8, 3), 52, horas=3, prep=None, traslado="")])
    assert leer_horas(libro, date(2026, 8, 1))[0]["horas_total"] == 3.0


def test_el_id_de_maquina_cruza_venga_como_texto_o_como_numero():
    """En el Sheet publicado el id viene como texto y en una copia bajada a
    mano como número: si no se normaliza, el cruce no encuentra nada."""
    libro = _libro(
        [_hora(date(2026, 8, 3), "52"), _hora(date(2026, 8, 4), 735.0)],
        maestro=[("52", "TRACTOR A55 DEUTZ N°0001 TATO", "TRACTORES", "TERCEROS"),
                 (735.0, "CARGADORA MANITOU N°9872 PABLO ROJAS", "CARGADORAS", "TERCEROS")],
    )
    filas = leer_horas(libro, date(2026, 8, 1))
    assert [f["id_maquina"] for f in filas] == [52, 735]


def test_la_propiedad_se_compara_sin_espacios_ni_mayusculas():
    libro = _libro(
        [_hora(date(2026, 8, 3), 52)],
        maestro=[(52, "TRACTOR A55 DEUTZ N°0001 TATO", "TRACTORES", " terceros ")],
    )
    assert len(leer_horas(libro, date(2026, 8, 1))) == 1


# ─── De qué tercero es la máquina ───────────────────────────────────────────

@pytest.mark.parametrize("nombre,tercero", [
    ("TRACTOR A55 DEUTZ N°0001 TATO", "TATO"),
    ("CARGADORA MANITOU N°9872 PABLO ROJAS", "PABLO ROJAS"),
    ("MAQUINARIA RUBEN MUÑOZ", "RUBEN MUÑOZ"),
    ("CINCEL SAN MIGUEL", "CINCEL SAN MIGUEL"),   # sin número ni prefijo: el nombre entero
    ("", ""),
    (None, ""),
])
def test_el_tercero_se_deduce_del_nombre_de_la_maquina(nombre, tercero):
    assert _detectar_tercero(nombre) == tercero


def test_el_numero_de_interno_separado_del_simbolo_rompe_la_deteccion():
    """Caso real de agosto 2026: 'N°' y '113' quedan como palabras distintas,
    ninguna arranca con N y trae dígitos a la vez, así que no hay corte y el
    tercero sale con el nombre entero de la máquina. Está fijado como está
    —el Excel hace exactamente lo mismo— para que la etapa 1 coincida con él;
    la solución no es afinar la heurística sino que el sistema de campo traiga
    el dueño como campo propio (plan-terceros.md, sección 2.1)."""
    assert _detectar_tercero("TRACTOR DEUTZ N° 113 SOSA ALBERTO") == (
        "TRACTOR DEUTZ N° 113 SOSA ALBERTO"
    )


def test_cuando_hay_dos_numeros_corta_por_el_ultimo():
    assert _detectar_tercero("TRACTOR N°15 ACOPLADO N°22 ANGEL ROJAS") == "ANGEL ROJAS"


# ─── Que falle a la vista ───────────────────────────────────────────────────

def test_si_el_sheet_cambia_de_formato_lo_dice():
    """No adivinar por posición: si falta una columna, se avisa."""
    wb = openpyxl.Workbook()
    hm = wb.active
    hm.title = consulta_taller.HOJA_MAESTRO
    hm.append(consulta_taller.COLUMNAS_MAESTRO)
    hh = wb.create_sheet(consulta_taller.HOJA_HORAS)
    hh.append(("Fecha", "Horas"))          # le faltan las demás
    buffer = io.BytesIO()
    wb.save(buffer)

    with pytest.raises(ValueError) as error:
        leer_horas(buffer.getvalue())
    assert "Horas_Preparacion" in str(error.value)


def test_sin_url_configurada_avisa_en_vez_de_inventar_un_origen(monkeypatch):
    monkeypatch.setattr(consulta_taller.settings, "taller_sheet_url", "")
    with pytest.raises(TallerNoConfigurado):
        descargar_libro()


def test_el_servicio_usa_el_libro_que_se_le_pasa_y_no_sale_a_la_red(monkeypatch):
    def no_llamar(*a, **k):
        raise AssertionError("no debería bajar el Sheet si ya le dieron el libro")
    monkeypatch.setattr(consulta_taller, "descargar_libro", no_llamar)

    libro = _libro([_hora(date(2026, 8, 3), 52)])
    filas = ConsultaTallerService(libro=libro).horas_quincena(date(2026, 8, 1))
    assert len(filas) == 1
